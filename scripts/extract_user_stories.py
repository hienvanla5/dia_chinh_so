import os
import sys
import openpyxl

# Dam bao output khong bi loi encoding tren Windows console
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def extract_all():
    excel_path = 'Địa Chính Số.xlsx'
    output_dir = 'user_stories'
    os.makedirs(output_dir, exist_ok=True)

    print(f"Loading {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # 1. Load Business Rules
    rules_dict = {}
    ws_rules = wb['Business Rules (Quy tắc)']
    rules_rows = list(ws_rules.iter_rows(values_only=True))
    for row in rules_rows[1:]:
        if not row[0]:
            continue
        rule_id = str(row[0]).strip()
        rules_dict[rule_id] = {
            'id': rule_id,
            'name': str(row[1] or '').strip(),
            'desc': str(row[2] or '').strip(),
            'roles': str(row[3] or '').strip(),
            'data_object': str(row[4] or '').strip(),
            'condition': str(row[5] or '').strip(),
            'then_action': str(row[6] or '').strip(),
            'else_action': str(row[7] or '').strip(),
            'related_stories': str(row[8] or '').strip(),
        }

    # 2. Load Acceptance Criteria grouped by Story ID
    ac_dict = {}
    ws_ac = wb['Acceptance Criteria']
    ac_rows = list(ws_ac.iter_rows(values_only=True))
    for row in ac_rows[1:]:
        if not row[1]:
            continue
        story_id = str(row[1]).strip()
        if story_id not in ac_dict:
            ac_dict[story_id] = []
        ac_dict[story_id].append({
            'no': row[0],
            'ac_id': str(row[2] or '').strip(),
            'scenario_type': str(row[3] or '').strip(),
            'given': str(row[4] or '').strip(),
            'when': str(row[5] or '').strip(),
            'then': str(row[6] or '').strip(),
            'test_data': str(row[7] or '').strip(),
            'severity': str(row[8] or '').strip(),
        })

    # 3. Load Tasks grouped by Story ID
    task_dict = {}
    ws_tasks = wb['Tasks (Công việc)']
    task_rows = list(ws_tasks.iter_rows(values_only=True))
    for row in task_rows[1:]:
        if not row[1]:
            continue
        story_id = str(row[1]).strip()
        if story_id not in task_dict:
            task_dict[story_id] = []
        task_dict[story_id].append({
            'no': row[0],
            'task_id': str(row[2] or '').strip(),
            'task_name': str(row[3] or '').strip(),
            'expected_result': str(row[4] or '').strip(),
            'task_type': str(row[5] or '').strip(),
            'assignee': str(row[6] or '').strip(),
            'priority': str(row[7] or '').strip(),
            'points': str(row[8] or '').strip(),
            'sprint': str(row[9] or '').strip(),
            'status': str(row[10] or '').strip(),
            'notes': str(row[11] or '').strip(),
        })

    # 4. Load User Stories
    stories = []
    ws_us = wb['Product Backlog (User Stories)']
    us_rows = list(ws_us.iter_rows(values_only=True))
    for row in us_rows[1:]:
        if not row[3]:
            continue
        story_id = str(row[3]).strip()
        rule_codes_raw = str(row[11] or '').strip()
        rule_codes = [r.strip() for r in rule_codes_raw.replace(';', ',').split(',') if r.strip()]

        stories.append({
            'no': row[0],
            'epic_id': str(row[1] or '').strip(),
            'epic_name': str(row[2] or '').strip(),
            'story_id': story_id,
            'title': str(row[4] or '').strip(),
            'user_story': str(row[5] or '').strip(),
            'role': str(row[6] or '').strip(),
            'business_value': str(row[7] or '').strip(),
            'description': str(row[8] or '').strip(),
            'precondition': str(row[9] or '').strip(),
            'postcondition': str(row[10] or '').strip(),
            'rule_codes': rule_codes,
            'priority': str(row[12] or '').strip(),
            'points': str(row[13] or '').strip(),
            'dependency': str(row[14] or '').strip(),
            'assignee': str(row[15] or '').strip(),
            'status': str(row[16] or '').strip(),
            'notes': str(row[17] or '').strip(),
        })

    print(f"Processing {len(stories)} stories, {len(rules_dict)} rules, {sum(len(v) for v in ac_dict.values())} ACs, {sum(len(v) for v in task_dict.values())} tasks...")

    # 5. Generate individual story markdown files
    for s in stories:
        sid = s['story_id']
        filename = f"{sid}.md"
        filepath = os.path.join(output_dir, filename)

        acs = ac_dict.get(sid, [])
        tasks = task_dict.get(sid, [])

        # Format AC table
        ac_lines = []
        if acs:
            ac_lines.append("| Mã AC | Loại kịch bản | Điều kiện ban đầu (Given) | Hành động (When) | Kết quả mong đợi (Then) | Dữ liệu kiểm thử | Mức độ |")
            ac_lines.append("|---|---|---|---|---|---|---|")
            for ac in acs:
                ac_lines.append(f"| **{ac['ac_id']}** | {ac['scenario_type']} | {ac['given']} | {ac['when']} | {ac['then']} | `{ac['test_data']}` | {ac['severity']} |")
        else:
            ac_lines.append("*Chưa có tiêu chí chấp nhận cụ thể.*")

        # Format Tasks checklist
        task_lines = []
        if tasks:
            for t in tasks:
                task_lines.append(f"- [ ] **{t['task_id']}**: {t['task_name']}")
                task_lines.append(f"  - *Kết quả cần đạt*: {t['expected_result']}")
                task_lines.append(f"  - *Nhóm CV*: `{t['task_type']}` | *Phụ trách*: {t['assignee']} | *Độ ưu tiên*: {t['priority']} | *Points*: {t['points']}")
        else:
            task_lines.append("*Chưa có danh sách công việc chi tiết.*")

        # Format Business Rules
        rule_lines = []
        if s['rule_codes']:
            for rcode in s['rule_codes']:
                if rcode in rules_dict:
                    r = rules_dict[rcode]
                    rule_lines.append(f"### [{r['id']}] {r['name']}")
                    rule_lines.append(f"- **Mô tả**: {r['desc']}")
                    rule_lines.append(f"- **Vai trò áp dụng**: {r['roles']}")
                    rule_lines.append(f"- **Đối tượng dữ liệu**: `{r['data_object']}`")
                    rule_lines.append(f"- **Điều kiện (Condition)**: {r['condition']}")
                    rule_lines.append(f"- **Kết quả khi thỏa (Then)**: {r['then_action']}")
                    rule_lines.append(f"- **Xử lý vi phạm (Else)**: {r['else_action']}")
                    rule_lines.append("")
                else:
                    rule_lines.append(f"- **Quy tắc**: `{rcode}`")
        else:
            rule_lines.append("*Không có quy tắc nghiệp vụ đặc thù riêng.*")

        # Format Dependencies & Related stories
        dep_str = s['dependency'] if s['dependency'] and s['dependency'] != 'None' and s['dependency'] != '-' else 'Không có (Story độc lập)'
        dep_links = []
        if dep_str != 'Không có (Story độc lập)':
            for dep_id in dep_str.replace(';', ',').split(','):
                dep_id = dep_id.strip()
                if dep_id:
                    dep_links.append(f"[{dep_id}](./{dep_id}.md)")
            dep_display = ", ".join(dep_links) if dep_links else dep_str
        else:
            dep_display = dep_str

        content = f"""# [{s['story_id']}] {s['title']}

## 1. Thông Tin Chung (Metadata)
- **Mã User Story**: `{s['story_id']}`
- **Thuộc Epic**: `{s['epic_id']}` - {s['epic_name']}
- **Vai trò chính (Persona)**: {s['role']}
- **Độ ưu tiên**: {s['priority']}
- **Story Points ước tính**: {s['points']}
- **Trạng thái**: {s['status']}
- **Người phụ trách đề xuất**: {s['assignee']}

---

## 2. Mô Tả User Story
> **Là một**: {s['role']},  
> **Tôi muốn**: {s['user_story'].split(', tôi muốn ')[1].split(', để ')[0] if ', tôi muốn ' in s['user_story'] and ', để ' in s['user_story'] else s['user_story']},  
> **Để**: {s['user_story'].split(', để ')[1] if ', để ' in s['user_story'] else s['business_value']}

- **Giá trị nghiệp vụ**: {s['business_value']}
- **Mô tả chi tiết giải pháp**: {s['description']}
- **Điều kiện tiên quyết (Precondition)**: {s['precondition']}
- **Kết quả mong đợi sau hoàn thành (Postcondition)**: {s['postcondition']}

---

## 3. Tiêu Chí Chấp Nhận (Acceptance Criteria - {len(acs)} tiêu chí)
{chr(10).join(ac_lines)}

---

## 4. Danh Sách Công Việc Kỹ Thuật (Tasks - {len(tasks)} tasks)
{chr(10).join(task_lines)}

---

## 5. Quy Tắc Nghiệp Vụ Ràng Buộc (Business Rules)
{chr(10).join(rule_lines)}

---

## 6. Phụ Thuộc & User Stories Liên Quan
- **Phụ thuộc vào**: {dep_display}
- **Ghi chú kỹ thuật**: {s['notes'] or 'Không có ghi chú thêm.'}

---

## 7. Kế Hoạch Kiểm Thử Tự Động (Automated Testing Plan)
- [ ] **Unit Tests**:
  - Test validation đầu vào với Zod Schema tương ứng.
  - Test service xử lý logic nghiệp vụ và ngoại lệ theo từng kịch bản AC.
- [ ] **Integration Tests**:
  - Kiểm thử Controller Endpoint (HTTP status, Response payload, Error codes).
  - Kiểm thử tương tác cơ sở dữ liệu (PostGIS, Drizzle ORM) và Redis Cache/Lock (nếu có).
- [ ] **Verification Gate**:
  - Chạy `pnpm test` đạt 100% pass.
  - Chạy `pnpm build` không có lỗi Type hoặc Build.
"""

        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(content)

    print(f"Successfully generated {len(stories)} user story markdown files in '{output_dir}/'.")

    # 6. Generate Master README.md index in user_stories/
    readme_path = os.path.join(output_dir, 'README.md')
    readme_lines = [
        "# Danh Mục Toàn Bộ User Stories - Hệ Thống Địa Chính Số",
        "",
        "> Tổng hợp toàn bộ 114 User Stories thuộc 12 Nhóm chức năng (Epics) được số hóa từ tài liệu chuẩn `Địa Chính Số.xlsx`.",
        "",
        "## Mục Lục Các Nhóm Chức Năng (Epics)",
        ""
    ]

    # Group stories by Epic
    epics = {}
    for s in stories:
        eid = s['epic_id']
        if eid not in epics:
            epics[eid] = {
                'id': eid,
                'name': s['epic_name'],
                'stories': []
            }
        epics[eid]['stories'].append(s)

    for eid, e in epics.items():
        readme_lines.append(f"### Epic {e['id']}: {e['name']} ({len(e['stories'])} User Stories)")
        readme_lines.append("")
        readme_lines.append("| STT | Mã Story | Tên User Story | Vai trò | Điểm | Độ ưu tiên | Trạng thái |")
        readme_lines.append("|---|---|---|---|---|---|---|")
        for s in e['stories']:
            sid = s['story_id']
            readme_lines.append(f"| {s['no']} | [{sid}](./{sid}.md) | [{s['title']}](./{sid}.md) | {s['role']} | {s['points']} | {s['priority']} | {s['status']} |")
        readme_lines.append("")

    with open(readme_path, 'w', encoding='utf-8') as f:
        f.write("\n".join(readme_lines))

    print(f"Successfully generated master index file at '{readme_path}'.")

if __name__ == '__main__':
    extract_all()
